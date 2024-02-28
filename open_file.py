from tkinter import filedialog, messagebox, ttk
import pandas as pd
import openpyxl
from docx import Document
import tkinter as tk



def treeview(treeview_frame):
    treeview = ttk.Treeview(treeview_frame, show="headings", style="Treeview")
    treeview["columns"] = tuple()
    treeview.pack(expand=True, fill=tk.BOTH)
def display_in_treeview(treeview, df):
    # Clear existing data
    for item in treeview.get_children():
        treeview.delete(item)

    for col in treeview["columns"]:
        treeview.heading(col, text="")
        treeview.column(col, width=0)

    treeview["columns"] = tuple(['Row No.'] + list(df.columns))

    for col in treeview["columns"]:
        treeview.heading(col, text=col)
        treeview.column(col, width=100)

    for index, row in df.iterrows():
        treeview.insert("", index, values=tuple([index + 1] + list(row)))

def open_file(file_type):
    file_extension = file_type.lower()
    file_path = filedialog.askopenfilename(title=f"Select {file_type} File", filetypes=[(f"{file_type} files", f"*.{file_extension}")])
    if file_path:
        print(f"Opening {file_type} file: {file_path}")
        try:
            display_data(file_path, file_type)
        except Exception as e:
            messagebox.showerror("Error", f"Error reading {file_type} file: {e}")

def display_data(file_path, file_type):
    if file_type.lower() == 'csv':
        current_data = pd.read_csv(file_path)
    elif file_type.lower() == 'text':
        with open(file_path, 'r') as file:
            text_data = file.read()
        current_data = pd.DataFrame({"Text Data": [text_data]})
    elif file_type.lower() == 'excel':
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        data = []
        for row in sheet.iter_rows(min_row=1, values_only=True):
            data.append(row)
        workbook.close()
        current_data = pd.DataFrame(data, columns=sheet[1])
    elif file_type.lower() == 'word':
        document = Document(file_path)
        text_data = ""
        for paragraph in document.paragraphs:
            text_data += paragraph.text + "\n"
        current_data = pd.DataFrame({"Text Data": [text_data]})
    

    display_in_treeview(treeview,current_data,)
    
    file_path = file_path